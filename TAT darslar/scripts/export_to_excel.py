# -*- coding: utf-8 -*-
import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

template_path = r"c:\Users\user\Desktop\TAT darslar\namun test importga.xlsx"
json_path = r"c:\Users\user\Desktop\TAT darslar\TAT darslar\testlar_banki_200.json"

output_paths = [
    r"c:\Users\user\Desktop\TAT darslar\testlar_import_200.xlsx",
    r"c:\Users\user\Desktop\TAT darslar\TAT darslar\testlar_import_200.xlsx"
]

with open(json_path, "r", encoding="utf-8") as f:
    questions = json.load(f)

print(f"Yuklangan savollar soni: {len(questions)}")

# Shablonni ochamiz
wb = openpyxl.load_workbook(template_path)

if "Savollar" not in wb.sheetnames:
    raise ValueError("Shablonda 'Savollar' varag'i topilmadi!")

ws = wb["Savollar"]

# Eski namuna qatorlarini (2-qatordan boshlab) tozalaymiz
max_r = ws.max_row
if max_r >= 2:
    ws.delete_rows(2, max_r - 1)

letters = ["A", "B", "C", "D"]

# Stillar
font_body = Font(name="Calibri", size=11, bold=False)
font_bold = Font(name="Calibri", size=11, bold=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_center = Alignment(horizontal="center", vertical="center")
thin_border = Border(
    left=Side(style='thin', color='FFD1D5DB'),
    right=Side(style='thin', color='FFD1D5DB'),
    top=Side(style='thin', color='FFD1D5DB'),
    bottom=Side(style='thin', color='FFD1D5DB')
)

# Savollarni kiritamiz
for idx, q in enumerate(questions, start=1):
    row_num = idx + 1
    q_num = idx
    q_text = q["q"]
    opt_a = q["opts"][0]
    opt_b = q["opts"][1]
    opt_c = q["opts"][2]
    opt_d = q["opts"][3]
    correct_ans = letters[q["ans"]]

    ws.cell(row=row_num, column=1, value=q_num).alignment = align_center
    ws.cell(row=row_num, column=2, value=q_text).alignment = align_left
    ws.cell(row=row_num, column=3, value=opt_a).alignment = align_left
    ws.cell(row=row_num, column=4, value=opt_b).alignment = align_left
    ws.cell(row=row_num, column=5, value=opt_c).alignment = align_left
    ws.cell(row=row_num, column=6, value=opt_d).alignment = align_left
    ws.cell(row=row_num, column=7, value=correct_ans).alignment = align_center

    # Shrift va chegaralar
    ws.cell(row=row_num, column=1).font = font_bold
    ws.cell(row=row_num, column=7).font = font_bold
    for col_idx in range(1, 8):
        c = ws.cell(row=row_num, column=col_idx)
        c.border = thin_border
        if col_idx in (2, 3, 4, 5, 6):
            c.font = font_body

# Ustun kengliklarini moslaymiz
ws.column_dimensions['A'].width = 6.0
ws.column_dimensions['B'].width = 60.0
ws.column_dimensions['C'].width = 25.0
ws.column_dimensions['D'].width = 25.0
ws.column_dimensions['E'].width = 25.0
ws.column_dimensions['F'].width = 25.0
ws.column_dimensions['G'].width = 16.0

for out_path in output_paths:
    wb.save(out_path)
    print(f"Muvaffaqiyatli saqlandi: {out_path}")

print("Tayyor! Jami 200 ta savol namunadagidek Excel formatiga o'tkazildi.")
